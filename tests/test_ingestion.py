"""
Tests unitarios e integrados para el Pipeline de Ingesta (Fase 6).
Verifica:
1. Lectores modulares (PDF, DOCX).
2. Detección de falta de capa de texto en PDF.
3. Rechazo explícito de .doc y formatos no soportados.
4. Segmentador lógico de documentos para FTS5.
5. Cálculo de hash SHA-256 y detección de cambios.
6. Pipeline completo con FakeLLM, idempotencia y aislamiento de fallos.
"""

import docx
import pypdf
import pytest

from procesa_agent.core.errors import IngestionError
from procesa_agent.domain.models import FichaProyecto, Iniciativa, LeccionAprendida, MetricaKPI
from procesa_agent.infrastructure.llm.fake import FakeLLM
from procesa_agent.ingestion.pipeline import IngestionPipeline
from procesa_agent.ingestion.readers import DocxReader, PDFReader, get_reader
from procesa_agent.ingestion.segmenter import DocumentSegmenter


def test_get_reader_formatos_soportados_y_rechazados(tmp_path):
    """Verifica que la fábrica get_reader seleccione el lector o lance IngestionError."""
    pdf_file = tmp_path / "test.pdf"
    docx_file = tmp_path / "test.docx"
    doc_file = tmp_path / "test.doc"
    txt_file = tmp_path / "test.txt"

    assert isinstance(get_reader(pdf_file), PDFReader)
    assert isinstance(get_reader(docx_file), DocxReader)

    # .doc debe ser rechazado con mensaje orientador
    with pytest.raises(IngestionError) as exc_doc:
        get_reader(doc_file)
    assert "formato heredado '.doc' no soportado" in str(exc_doc.value).lower()

    # Formatos desconocidos
    with pytest.raises(IngestionError) as exc_txt:
        get_reader(txt_file)
    assert "formato no soportado" in str(exc_txt.value).lower()


def test_pdf_reader_sin_capa_de_texto(tmp_path):
    """Verifica que PDFReader detecte PDFs sin capa de texto y lance IngestionError."""
    vacio_pdf = tmp_path / "vacio.pdf"
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with open(vacio_pdf, "wb") as f:
        writer.write(f)

    reader = PDFReader()
    with pytest.raises(IngestionError) as exc:
        reader.read(vacio_pdf)
    assert "sin capa de texto" in str(exc.value).lower()


def test_docx_reader_extraccion_parrafos_y_tablas(tmp_path):
    """Verifica que DocxReader extraiga párrafos y tablas preservando contenido."""
    docx_file = tmp_path / "informe_test.docx"
    doc = docx.Document()
    doc.add_paragraph("1. Resumen ejecutivo")
    doc.add_paragraph("Este es el contenido del resumen ejecutivo.")

    # Agregar tabla
    tabla = doc.add_table(rows=2, cols=2)
    tabla.cell(0, 0).text = "KPI"
    tabla.cell(0, 1).text = "Resultado"
    tabla.cell(1, 0).text = "Eficiencia"
    tabla.cell(1, 1).text = "95%"

    doc.save(str(docx_file))

    reader = DocxReader()
    res = reader.read(docx_file)
    assert res["archivo_origen"] == "informe_test.docx"
    assert "Resumen ejecutivo" in res["texto_completo"]
    assert "Eficiencia | 95%" in res["texto_completo"]


def test_segmenter_logico():
    """Verifica que DocumentSegmenter divida en secciones tituladas."""
    texto = (
        "Metadatos iniciales del proyecto de consultoría.\n\n"
        "1. Diagnóstico Inicial\n"
        "Se detectaron problemas en la línea de producción.\n\n"
        "2. Metodología Aplicada\n"
        "Se aplicó Lean Six Sigma y SMED.\n\n"
        "Anexo A. Detalle de Tiempos\n"
        "Cronograma detallado de actividades."
    )

    segmenter = DocumentSegmenter()
    secciones = segmenter.segment(texto)

    assert len(secciones) == 4
    titulos = [s[0] for s in secciones]
    assert "Cabecera y Metadatos" in titulos
    assert "1. Diagnóstico Inicial" in titulos
    assert "2. Metodología Aplicada" in titulos
    assert "Anexo A. Detalle de Tiempos" in titulos


def test_pipeline_ingesta_con_fake_llm_e_idempotencia_sha(tmp_path):
    """Verifica ejecución del pipeline completo con FakeLLM, hash SHA-256 e idempotencia."""
    db_file = str(tmp_path / "test_pipeline.sqlite")
    raw_dir = tmp_path / "raw"
    fichas_dir = tmp_path / "fichas"
    raw_dir.mkdir()
    fichas_dir.mkdir()

    # Crear documento DOCX de prueba
    doc_path = raw_dir / "Informe_Test_001.docx"
    doc = docx.Document()
    doc.add_paragraph("1. Resumen Ejecutivo")
    doc.add_paragraph("Proyecto de optimización logística para Distribuidora Global.")
    doc.save(str(doc_path))

    ficha_mock = FichaProyecto(
        codigo_proyecto="TEST-LOG-01",
        archivo_origen="Informe_Test_001.docx",
        cliente="Distribuidora Global S.A.",
        sector="Logística",
        periodo="2025",
        duracion_semanas=8,
        gerente_proyecto="Carlos Ruiz",
        estado="Cerrado aceptado",
        resumen_ejecutivo="Optimización logística completada.",
        iniciativas_clave=[
            Iniciativa(
                nombre="Ruteo Dinámico",
                descripcion="Rutas inteligentes",
                impacto_esperado="Menor costo",
            )
        ],
        kpis=[
            MetricaKPI(
                indicador="Tiempos de Entrega",
                unidad="h",
                linea_base="48 h",
                meta="< 24 h",
                resultado="20 h",
                variacion="-58%",
                cumplimiento="Cumplido",
                observaciones="Objetivo superado",
            )
        ],
        lecciones=[
            LeccionAprendida(
                tema="Metodología", titulo="Adopción de software", descripcion="Capacitación clave"
            )
        ],
        proximos_pasos=["Monitoreo continuo"],
    )

    fake_llm = FakeLLM(structured_response=ficha_mock)

    pipeline = IngestionPipeline(
        db_path=db_file,
        raw_dir=raw_dir,
        fichas_dir=fichas_dir,
        llm_provider=fake_llm,
    )

    # 1. Primera ejecución: debe procesar el archivo
    reporte_1 = pipeline.ejecutar()
    assert reporte_1["total_archivos"] == 1
    assert "Informe_Test_001.docx" in reporte_1["procesados"]
    assert reporte_1["omitidos"] == []
    assert reporte_1["fallidos"] == []
    assert (fichas_dir / "TEST-LOG-01.json").exists()

    # 2. Segunda ejecución sin cambios: debe omitir por hash SHA-256
    reporte_2 = pipeline.ejecutar(force=False)
    assert "Informe_Test_001.docx" in reporte_2["omitidos"]
    assert reporte_2["procesados"] == []

    # 3. Tercera ejecución con force=True: debe re-procesar
    reporte_3 = pipeline.ejecutar(force=True)
    assert "Informe_Test_001.docx" in reporte_3["procesados"]


def test_pipeline_aislamiento_de_fallos(tmp_path):
    """Verifica que un archivo inválido no aborte el lote y se reporte en 'fallidos'."""
    db_file = str(tmp_path / "test_fallos.sqlite")
    raw_dir = tmp_path / "raw"
    fichas_dir = tmp_path / "fichas"
    raw_dir.mkdir()
    fichas_dir.mkdir()

    # Archivo inválido (.doc no soportado)
    doc_invalido = raw_dir / "archivo_antiguo.doc"
    doc_invalido.write_text("dummy", encoding="utf-8")

    pipeline = IngestionPipeline(
        db_path=db_file,
        raw_dir=raw_dir,
        fichas_dir=fichas_dir,
    )

    reporte = pipeline.ejecutar()
    assert len(reporte["fallidos"]) == 1
    assert reporte["fallidos"][0]["archivo"] == "archivo_antiguo.doc"
    assert "no soportado" in reporte["fallidos"][0]["error"].lower()
    assert reporte["exitoso"] is False
